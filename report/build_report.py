"""Build the midterm report (report/midterm_report.docx + .pdf).

Numbers come from the data (data/raw) and from report/results.json (experiment summary).
Run from the midterm/ folder:  python report/build_report.py
"""
import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
FIG = ROOT / 'report' / 'figures'
RES = json.loads((ROOT / 'report' / 'results.json').read_text(encoding='utf-8'))
CLASSES = ['Car', 'Motorcycle', 'Bus', 'Truck', 'Tuktuk', 'Van', 'Pickup', 'Songthaew']

# ---------------------------------------------------------------- data stats
raw = ROOT / 'data' / 'raw'
tr = pd.read_csv(raw / 'train.csv')
tr['cam'] = tr.image_id.str.split('_').str[0]
n_train_img = len(list((raw / 'train' / 'train').glob('*.jpg')))
n_test_img = len(list((raw / 'test' / 'test').glob('*.jpg')))
test_cams = sorted({p.name.split('_')[0] for p in (raw / 'test' / 'test').glob('*.jpg')})
train_cams = sorted({p.name.split('_')[0] for p in (raw / 'train' / 'train').glob('*.jpg')})
cls_boxes = tr.class_id.value_counts().sort_index()
cls_imgs = tr.groupby('class_id').image_id.nunique()
bw, bh = tr.x2 - tr.x1, tr.y2 - tr.y1

# ---------------------------------------------------------------- helpers
doc = Document()
sec = doc.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
for side in ('left_margin', 'right_margin'):
    setattr(sec, side, Cm(2.2))
sec.top_margin = sec.bottom_margin = Cm(2.0)
st = doc.styles['Normal']
st.font.name = 'Calibri'
st.font.size = Pt(11)
st.paragraph_format.space_after = Pt(6)
st.paragraph_format.line_spacing = 1.15

fig_no = [0]
tab_no = [0]


def H(text, level=1):
    return doc.add_heading(text, level=level)


def P(text='', bold=False, italic=False, align=None, size=None):
    p = doc.add_paragraph()
    # **bold** segments inside text
    parts = text.split('**')
    for i, part in enumerate(parts):
        r = p.add_run(part)
        r.bold = bold or (i % 2 == 1)
        r.italic = italic
        if size:
            r.font.size = Pt(size)
    if align:
        p.alignment = align
    return p


def B(items, style='List Bullet'):
    for it in items:
        p = doc.add_paragraph(style=style)
        parts = it.split('**')
        for i, part in enumerate(parts):
            p.add_run(part).bold = i % 2 == 1


def caption(prefix, counter, text):
    counter[0] += 1
    p = doc.add_paragraph()
    r = p.add_run(f'{prefix} {counter[0]}: {text}')
    r.italic = True
    r.font.size = Pt(9.5)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return counter[0]


def FIGURE(fn, text, width=16):
    path = FIG / fn
    if not path.exists():
        P(f'[missing figure: {fn}]', italic=True)
        return
    doc.add_picture(str(path), width=Cm(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption('Figure', fig_no, text)


def shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)


def TABLE(header, rows, text, widths=None, bold_rows=()):
    caption('Table', tab_no, text)
    t = doc.add_table(rows=1, cols=len(header))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ''
        r = c.paragraphs[0].add_run(str(h))
        r.bold = True
        r.font.size = Pt(9.5)
        shade(c, 'D9E2F3')
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ''
            r = cells[i].paragraphs[0].add_run(str(v))
            r.font.size = Pt(9.5)
            r.bold = ri in bold_rows
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


best = RES['experiments'][RES['best']]
final = RES.get('final', {})
fmt = lambda x: '-' if x is None else f'{x:.4f}'

# ---------------------------------------------------------------- title page
for _ in range(6):
    doc.add_paragraph()
P('2110531 Data Science and Data Engineering Tools (2026/1)', align=WD_ALIGN_PARAGRAPH.CENTER, size=13)
P('Take-Home Midterm Exam Report', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=22)
P('Bangkok CCTV Vehicle Detection', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=18)
doc.add_paragraph()
P(f"Name: {RES['student']['name']}", align=WD_ALIGN_PARAGRAPH.CENTER, size=12)
P(f"Student ID: {RES['student']['id']}", align=WD_ALIGN_PARAGRAPH.CENTER, size=12)
P(f"Kaggle username: {RES['student']['kaggle']}", align=WD_ALIGN_PARAGRAPH.CENTER, size=12)
doc.add_paragraph()
P(f"Source code, CSVs, experiment outputs (GitHub): {RES['links']['code']}", align=WD_ALIGN_PARAGRAPH.CENTER, size=10)
P(f"Model weights and prepared data (Google Drive): {RES['links']['weights']}", align=WD_ALIGN_PARAGRAPH.CENTER, size=10)
page_break()

# ---------------------------------------------------------------- 1 introduction
H('Chapter 1: Introduction')
P('The task is multi-class object detection of vehicles in Bangkok traffic-camera (CCTV) images. '
  'The images were scraped automatically every 30 minutes from the Bangkok Metropolitan Administration (BMA) '
  'traffic platform (cpudapp.bangkok.go.th/bmatraffic) between 25/08/2026 06:00 and 31/08/2026 21:00. '
  'After a quality screening, 20 cameras were kept. Every vehicle has to be localised with a bounding box and '
  'assigned to one of eight classes: Car, Motorcycle, Bus, Truck, Tuktuk, Van, Pickup and Songthaew.')
P('Submissions are scored on Kaggle with **mAP@50**: the COCO-style average precision at IoU 0.50, '
  'computed with pycocotools for each class and averaged over the eight classes. Because every class has the '
  'same weight, rare classes such as Songthaew or Tuktuk affect the score as much as Car does.')
P('Objectives of this work:')
B(['explore and clean the data, and build a validation scheme that reflects the hidden test set;',
   'fine-tune a modern pretrained detector (YOLO26 from the Ultralytics framework) and evaluate it with the '
   'same metric as Kaggle;',
   'analyse the errors and describe how the model could be improved.'])
P('The code is in two Jupyter notebooks: vehicle_detection.ipynb (EDA, cleaning, split, YOLO26s experiments, '
  'evaluation, visualisation, submission; local GPU) and cv_kfold_kaggle.ipynb / cv_kfold_kaggle_v2.ipynb (final YOLO26m '
  'camera 5-fold cross-validation v1 / v2; Kaggle GPU), plus experiments/ensemble_wbf.py for the final 10-model ensemble. No generative-AI / VLM API was used, so there is no API cost. '
  'The final submission (an ensemble of 10 YOLO26m models) reaches **0.630 mAP@50** on the public leaderboard.')

# ---------------------------------------------------------------- 2 data
H('Chapter 2: Data Preparation')
P('The analysis follows the data-preparation steps taught in the course (Week 2 slides and Lab1-Lab4): '
  '(1) examine the data set, (2) understand / narrow down the features and check for leakage, '
  '(3) handle outliers and skewness, and (4) build the train / validation split with a fixed random seed.')

H('2.1 Examining the data set', 2)
TABLE(['Item', 'Value'], [
    ['Train images', f'{n_train_img:,} ({len(train_cams)} cameras)'],
    ['Test images', f'{n_test_img:,} ({len(test_cams)} cameras); 997 are listed in sample_submission.csv'],
    ['Image size', '352 x 288 pixels (all images)'],
    ['Annotated boxes (train)', f'{len(tr):,} in {tr.image_id.nunique():,} images; {n_train_img - tr.image_id.nunique()} images have no vehicle'],
    ['Boxes per image', f'mean {len(tr) / n_train_img:.1f}, max {tr.groupby("image_id").size().max()}'],
    ['Box size (w x h)', f'median {bw.median():.0f} x {bh.median():.0f} px, 25% of boxes are smaller than {bw.quantile(.25):.0f} x {bh.quantile(.25):.0f} px'],
    ['Submission format', 'id, image_id, class_id, confidence, x1, y1, x2, y2 (one row per box)'],
], 'Dataset summary', widths=[4.5, 12])
TABLE(['Column', 'Type', 'Description', 'Missing'], [
    ['id', 'int', 'row id - identifier only, not a feature', 0],
    ['image_id', 'str', '<camera>_<YYYYMMDD>_<HHMMSS>.jpg; camera, date, hour and day/night are derived from it', 0],
    ['class_id', 'int (nominal)', 'target class 0-7', 0],
    ['x1, y1', 'int', 'top-left corner of the box [px]', 0],
    ['x2, y2', 'int', 'bottom-right corner of the box [px]', 0],
], 'Data dictionary of train.csv (27,396 rows x 7 columns)', widths=[2.2, 2.6, 10, 1.6])
P('**Out-of-range values and miscodes.** No missing values; all class ids are in 0-7; no inverted boxes and no box '
  'outside the image; every image_id exists on disk and all images are 352 x 288. Found: 25 exact duplicate boxes, '
  '15 boxes with a side < 2 px, and 102 training images without any vehicle (kept as negative examples).')
H('Target distribution', 3)
TABLE(['ID', 'Class', 'Boxes', 'Share', 'Images containing class', 'Imbalance vs Car'],
      [[k, CLASSES[k], f'{cls_boxes[k]:,}', f'{cls_boxes[k] / len(tr):.2%}', f'{cls_imgs[k]:,}',
        f'{cls_boxes[0] / cls_boxes[k]:.1f}x'] for k in range(8)],
      'Class distribution in the training set', widths=[1.1, 2.8, 2, 2, 4, 3.2])
FIGURE('eda_class_dist.png', 'Boxes per class (log scale) and fraction of training images that contain each class')
P('The target is heavily imbalanced: Car is 64.8% of all boxes while Songthaew is 0.43% (152 times fewer). Because '
  'mAP@50 averages the AP of the eight classes equally, the rare classes matter as much as Car.')

H('2.2 Feature understanding, association with the target and leakage', 2)
P(f'**Camera.** Train uses cameras {", ".join(train_cams)}; test uses {", ".join(test_cams)}. '
  '**The two sets are disjoint**, so the model must generalise to unseen viewpoints.')
FIGURE('eda_camera.png', 'Images per camera (train vs test) and boxes per camera x class')
st = RES['eda_stats']
P(f"**Chi-square test of camera x class** (slide appendix): chi2 = {st['chi2_cam']:,}, dof = 98, p < 1e-300, "
  f"Cramer's V = {st['v_cam']:.3f}. H0 (same class distribution for every camera) is rejected: the vehicle mix depends "
  "strongly on the camera. The standardised residuals show, for example, that 91 of the 117 Songthaews come from camera "
  "1426 (residual +53), Pickups are over-represented at cameras 232 and 222 and Tuktuks at 244, 229 and 172.")
FIGURE('eda_chi2_resid.png', 'Standardised chi-square residuals, camera x class (red = more boxes than expected)', width=14)
P(f"**Time of day.** Frames cover 06:00-21:00 every 30 minutes. Street lights keep night frames bright, so night was "
  f"defined by clock time (after 18:30, about 17% of frames). The class mix also changes at night (chi-square "
  f"p = {st['p_night']}, Cramer's V = {st['v_night']:.3f}); e.g. Songthaew rises from 0.27% of boxes by day to "
  "1.53% at night and Tuktuk from 0.84% to 1.86%.")
FIGURE('eda_brightness.png', 'Mean brightness vs hour of day', width=12)
FIGURE('eda_time.png', 'Traffic density by hour and class share by day / night')
P(f"**Leakage check.** In Lab1, columns that leak information were removed. Here the leak is at the image level: "
  f"consecutive frames of the same camera differ by a median of {st['mad_same']} grey levels, versus {st['mad_cross']} "
  f"for frames of different cameras ({st['mad_ratio']}x more similar). A random train / validation split would put "
  "almost identical scenes on both sides and measure memorisation instead of generalisation. Therefore **whole cameras "
  "are held out** for validation (Section 2.6).")
FIGURE('eda_leakage.png', 'Frame similarity: same camera (consecutive frames) vs different cameras', width=13)
P('**Train vs test distribution.** The test cameras have a similar share of night frames (17.6% vs 17.1%) and a '
  'slightly darker mean brightness (123 vs 132).')
FIGURE('eda_train_vs_test.png', 'Brightness and hour-of-day distributions of train and test images')

H('2.3 Outliers and skewness', 2)
P(f"Box sizes are tiny and strongly right-skewed: the box area has skew {st['skew_area']} and kurtosis "
  f"{st['kurt_area']}; after a log transform the skew is {st['skew_log']} (Lab4 approach). This motivates upscaling "
  "the images (imgsz 640 -> 960) and a multi-scale detector.")
FIGURE('eda_box_size.png', 'Boxes per image, box size per class and box width vs height')
FIGURE('eda_area_skew.png', 'Box area before and after log transform')
TABLE(['Class', 'Boxes', 'Median area', 'log-area IQR outliers', 'Median h/w', 'h/w IQR outliers'],
      RES['outlier_table'], 'Outliers per class (1.5 x IQR rule)', widths=[2.8, 1.8, 2.5, 3.4, 2.4, 3])
P('The extremes were inspected visually before deciding what to do. The largest boxes are real vehicles very close '
  'to camera 1426 (6 boxes are almost full height) and are kept. The extreme aspect ratios are annotation errors: '
  'slivers such as 1 x 8 or 2 x 16 px with no vehicle inside (195 boxes have a side <= 2 px). Boxes with a side '
  '< 2 px are removed; truncating coordinates, as done for numeric features in Lab4, is not meaningful for boxes.')
FIGURE('eda_extreme_boxes.png', 'Most extreme boxes: largest area (top), tallest (middle), widest (bottom)')

H('2.4 Visual inspection', 2)
FIGURE('eda_train_samples.png', 'Preview: one daytime frame per training camera with ground-truth boxes')
FIGURE('eda_train_night_samples.png', 'Preview: night frames of training cameras. Training uses all frames (day and night); only this preview is split by time')
FIGURE('eda_test_samples.png', 'Test cameras (day and night); these viewpoints never appear in training')
FIGURE('eda_crops.png', 'Examples of rare classes (crops). Van / Pickup / Truck are hard to separate even for a human', width=13)
P('**sample_submission.csv uses different ids from the test file names.** It lists 997 long ids '
  '(dataset_<cam>_annotated_coco1.0_<cam>_<Thai location name>_<date>_<time>.jpg), while the 1,013 test images are '
  'named <cam>_<date>_<time>.jpg. The Kaggle solution uses the file names (Section 4.3), so predictions are written '
  'for all 1,013 test images with the file name as image_id.')

H('2.5 Data cleaning', 2)
P('Cleaning is applied to the **training labels only**. Validation is always scored against the raw labels, '
  'the same way Kaggle scores the raw test labels.')
cl = RES['cleaning']
TABLE(['Step', 'Count'], [[k, v] for k, v in cl.items()], 'Data-cleaning log', widths=[11, 3])
P('280 pairs of boxes overlap with IoU > 0.9 but have different classes (mostly Car vs Truck and '
  'Bus vs Truck): the same vehicle was labelled twice. They were kept because the test labels probably '
  'contain the same pattern, and a model that outputs both classes at lower confidence can still match both.')

AU = json.loads((ROOT / 'report' / 'audit_results.json').read_text(encoding='utf-8'))
H('Label audit: rules (Lab5) and model-assisted check', 3)
P('The rule-based cleaning above only removes boxes that are certainly broken. To check the remaining labels, two '
  'audits were run (experiments/label_audit.py) and the flagged cases were inspected visually; nothing was changed '
  'in the training data afterwards, so the submitted models stay reproducible.')
P('**(a) Outlier rules from Lab5** (z-score > 3 and 1.5 x IQR on log aspect ratio and log area, per class):')
TABLE(['Class', 'z>3 aspect', 'z>3 area', 'IQR aspect', 'IQR area', 'side <= 2 px'],
      [[r[0]] + r[1:] for r in AU['rule_table']], 'Boxes flagged by the Lab5 outlier rules', widths=[2.8, 2.4, 2.2, 2.4, 2.2, 2.6])
FIGURE('audit_zscore_aspect.png', 'Boxes with |z(log aspect)| > 3: mostly vehicles cut by the image border (valid), a few slivers (errors)')
P('Visual inspection shows that most rule outliers are **valid**: vehicles cut by the image border (only the roof is '
  'visible) and thin, head-on motorcycles (163 of the 172 boxes with a side <= 3 px and ratio >= 4 are Motorcycles). '
  'Only a handful are clear annotation errors (e.g. 2 x 16 px or 7 x 1 px "Car" slivers). Removing every rule outlier, '
  'as for the numeric features of Lab5, would therefore delete hundreds of real vehicles, which also appear in the test set.')
o = AU['oof']
P('**(b) Model-assisted audit.** The out-of-fold predictions of the 5-fold CV are made by models that never saw the '
  'camera, so strong disagreements between model and label point to label problems:')
TABLE(['Check', 'Rule', 'Flagged', 'Visual review of the top/random 24'], [
    ['Possible missing label', 'prediction conf >= 0.7, IoU < 0.3 with every label', o['possible_missing_labels'],
     'about 20 of 24 are real unlabelled vehicles, mostly cut by the image border or very close to the camera'],
    ['Possible wrong class', 'best prediction (IoU >= 0.7) has another class, conf >= 0.6', o['possible_wrong_class'],
     'about 3 of 4 look like label errors: vans labelled Truck / Bus, pickups labelled Car'],
    ['Possible bad box', 'no prediction with IoU >= 0.3', o['possible_bad_boxes (no prediction)'],
     'mostly correct labels of tiny, distant vehicles in dense traffic (median 10 x 14 px): a model weakness, not a label error'],
], 'Model-assisted label audit (OOF predictions of YOLO26m)', widths=[3, 4.6, 1.6, 7.3])
P('Most frequent class disagreements (label -> model): ' + ', '.join(f'{k} ({v})' for k, v in list(o['wrong_pairs (label -> model)'].items())[:6]) +
  '. Part of the Truck -> Car count comes from the 212 Car/Truck double labels found in Section 2.5. '
  'Possible missing labels touch the image border in 23% of the cases, versus 9% of all labels: vehicles that are '
  'only partly visible were labelled inconsistently.')
FIGURE('audit_missing.png', 'Possible missing labels: confident predictions (red) where no label exists (green = existing labels)')
FIGURE('audit_wrong_class.png', 'Possible wrong classes: label (green) vs. confident model class')
FIGURE('audit_bad_boxes.png', 'Labels without any prediction (random sample): mostly tiny distant vehicles')
P('**Decision.** The labels contain noise of the same kind that the test labels probably contain (missing border '
  'vehicles, Van / Pickup / Truck ambiguity), so the training labels were not edited for the submitted models. The '
  'audit explains part of the remaining error (Chapter 5) and gives a concrete list for relabelling in future work '
  '(experiments/audit_*.csv).')

H('2.6 Train / validation split and dataset construction', 2)
P(f"To mimic the unseen-camera test set and avoid the leakage shown in 2.2, **{len(RES['val_cams'])} whole cameras "
  f"({', '.join(RES['val_cams'])}) were held out for validation**. As a stratification over cameras, the notebook "
  "tries all 1,365 combinations of 4 cameras out of 15 and keeps the one where every class has at least 10 validation "
  "boxes and the validation share of each class is closest to 4/15. The random seed is fixed (42) everywhere.")
TABLE(['Class', 'Train boxes', 'Val boxes', 'Val share', 'Train %', 'Val %'],
      [[c, *RES['split_table'][c]] for c in CLASSES],
      'Camera-held-out split after cleaning: 2,194 train / 797 val images (class proportions compared as in Lab3)',
      widths=[3, 2.4, 2.4, 2.2, 2.2, 2.2])
P('**Repeat-factor sampling (RFS).** mAP weights all classes equally, so images that contain rare classes '
  'are repeated during training (LVIS-style): for class c with image frequency f_c, r_c = max(1, sqrt(t / f_c)) '
  'with t = 0.15, and each image is repeated r = max_c r_c times with stochastic rounding. Songthaew images are '
  'repeated about 2.8x, Tuktuk 1.4x and Van 1.2x. Repeats are only added to the training list, never to validation.')
P('Labels were converted to the YOLO format (class x_center y_center width height, normalised) and '
  'the image lists were written as train.txt / val.txt / full.txt, which Ultralytics reads directly.')

H('2.7 Summary: findings and decisions', 2)
TABLE(['Finding', 'Decision'], [
    ['Test cameras never appear in training; same-camera frames are near-duplicates', 'Hold out whole cameras for validation'],
    ['Class mix depends on camera and time (chi-square p ~ 0)', 'Every class must be present in val; RFS for rare classes'],
    ['Car : Songthaew = 152 : 1', 'Repeat-factor sampling; per-class AP analysis'],
    ['Tiny, right-skewed box sizes', 'Upscale to imgsz 960; multi-scale detector'],
    ['Sliver boxes are label noise; duplicates', 'Remove side < 2 px and exact / near duplicates'],
    ['Lab5 outliers are mostly valid (border-cut vehicles, thin motorcycles)', 'Do not drop rule outliers; inspect visually'],
    ['OOF audit: missing border vehicles, Van/Pickup/Truck class noise', 'Keep labels (test has the same noise); list for relabelling'],
    ['17% night frames in both train and test', 'Report day / night mAP; HSV augmentation'],
], 'EDA findings and the resulting decisions', widths=[9, 7.5])

# ---------------------------------------------------------------- 3 model
CVR = json.loads((ROOT / 'report' / 'cv_results.json').read_text(encoding='utf-8'))
V12 = json.loads((ROOT / 'report' / 'v1_v2_oof.json').read_text(encoding='utf-8'))
V2F = [json.loads(Path(q).read_text()) for q in sorted((ROOT / 'experiments' / 'kaggle_cv_run_v2' / 'cv_out_v2').glob('fold*/metrics.json'))]
WBFT = pd.read_csv(ROOT / 'report' / 'wbf_tuning.csv')
H('Chapter 3: Model')
H('3.1 Choice of model', 2)
P('The course slides (Week 5, Object Detection) describe the change from two-stage CNNs (Faster R-CNN) to '
  'one-stage YOLO detectors, transformer detectors (DETR / RT-DETR) and, most recently, end-to-end NMS-free '
  'YOLO26. The course example notebook fine-tunes YOLOv8 with Ultralytics. This work uses the newest model in '
  'that family, **YOLO26**, fine-tuned from COCO-pretrained weights. Two sizes were used: **YOLO26s** for the '
  'experiments on the local 4 GB GPU, and **YOLO26m** for the final model, trained on Kaggle GPUs with camera '
  '5-fold cross-validation, twice: v1 (standard augmentation) and v2 (+ cross-camera copy-paste). The final '
  'prediction fuses the 10 fold models (v1 + v2).')
TABLE(['Property', 'YOLO26s (experiments)', 'YOLO26m (final)'], [
    ['Framework', 'Ultralytics 8.4.163, PyTorch 2.8.0', 'Ultralytics 8.4.163, PyTorch (Kaggle image)'],
    ['Parameters', '9.95 M', '21.8 M'],
    ['FLOPs', '22.8 G at 640', '75.0 G at 640 / 169.5 G at 960'],
    ['Pretraining', 'COCO 2017 detection (80 classes, ~118k images)', 'same'],
    ['Backbone', 'CSP-style CNN with C3k2 blocks, SPPF and C2PSA (attention) block', 'same design, wider / deeper'],
    ['Neck', 'PAN/FPN feature pyramid, strides 8 / 16 / 32 (small vehicles)', 'same'],
    ['Head', 'Anchor-free, end-to-end (one-to-one) head: no NMS, no DFL', 'same'],
    ['Transfer', 'head rows of car, motorcycle, bus, truck initialised from the COCO classes', 'same'],
], 'YOLO26 models used', widths=[2.8, 7, 6.7])
P('YOLO is a one-stage detector: a single CNN pass predicts, for every location of three feature maps, a box '
  'and class scores. YOLO26 trains with a one-to-one assignment so each object gets a single prediction, which '
  'removes the non-maximum-suppression step, and it drops the Distribution Focal Loss. According to the '
  'Ultralytics documentation it also uses progressive loss balancing (ProgLoss) and small-target-aware label '
  'assignment (STAL), which suits this dataset because most vehicles are tiny.')
P('No generative-AI / VLM API was used, so there is no API cost. Compute: the YOLO26s runs used a local laptop GPU; '
  f"the YOLO26m cross-validations used the free Kaggle GPU quota (2 x NVIDIA T4): v1 {sum(f['hours'] for f in CVR['folds']):.1f} h and "
  f"v2 {sum(f['hours'] for f in V2F):.1f} h for 5 folds each (plus a 4 h A/B test of v2 on two folds), including prediction.")

H('3.2 Training setup', 2)
TABLE(['Hyper-parameter', 'YOLO26s (experiments)', 'YOLO26m 5-fold (v1 and v2)'], [
    ['Input size', '640 and 960 (upscaled from 352 x 288, letterbox)', '960'],
    ['Epochs / early stopping', '100 (640) / 40 (960), patience 30, best-val checkpoint', '30 per fold, patience 10, best-val checkpoint of each fold'],
    ['Batch size', '8 (640) / 4 (960), 4 GB VRAM', '16 (8 per GPU, 2 GPUs, DDP)'],
    ['Validation', '4 held-out cameras (1066, 1427, 1437, 244)', 'StratifiedGroupKFold by camera: 5 folds x 3 cameras'],
    ['Optimizer', 'Ultralytics "auto" -> AdamW (lr 0.000833, momentum 0.9, wd 0.0005), cosine LR, warm-up 3 epochs', 'same'],
    ['Augmentation', 'Mosaic (off for the last 10 epochs), scale 0.5, translate 0.1, flip 0.5, HSV (h 0.015, s 0.7, v 0.4); applied on the fly, re-sampled every time an image is loaded (no rotation, shear, vertical flip, mixup or copy-paste)', 'v1: same; v2: same + offline copy-paste and photometric variants (Section 3.3)'],
    ['Class balance', 'Repeat-factor sampling (Section 2.6)', 'same'],
    ['Seed', '42, deterministic=True', 'same'],
    ['Hardware', 'RTX 3050 Ti Laptop GPU (4 GB), 32 GB RAM', 'Kaggle Notebook, 2 x NVIDIA T4 (16 GB each)'],
], 'Training configuration', widths=[3, 6.8, 6.7])
P('**Camera K-fold cross-validation.** The single 4-camera split of Section 2.6 wastes 4 cameras for training and '
  'gives a noisy score. For the final model the 15 training cameras were split with scikit-learn '
  'StratifiedGroupKFold (group = camera, stratified on the rarest class of each image) into 5 folds of 3 cameras. '
  'Each fold model is trained on 12 cameras and early-stopped on its own 3 unseen cameras, so every camera is '
  'used for validation exactly once (out-of-fold score) and for training in 4 of the 5 models. The folds are '
  'fixed in the notebook because a different scikit-learn version produces a different assignment.')
H('3.3 v2: cross-camera copy-paste of rare classes', 2)
P('The CV of v1 showed that rare classes do not transfer between cameras (Songthaew from camera 1426 is not '
  'detected on camera 1066). v2 therefore adds training data built **per fold from the training cameras only**:')
B(['**Cross-camera copy-paste:** a crop of a Bus, Tuktuk, Van, Pickup or Songthaew from one camera is pasted into an '
   'image of a different camera, in the same lane as (or next to) an existing vehicle, scaled to that vehicle '
   '(x the typical size ratio of the two classes), at most 1.6x upscaling, never overlapping a labelled box, with a '
   'feathered edge; each source crop is used at most 3 times. 500 synthetic images per fold (~600 pasted vehicles).',
   '**Photometric variants** (lower contrast, gamma, light blur, JPEG noise) for the repeat-factor copies and the '
   'synthetic images, because the test cameras are darker and have lower contrast (Section 2.2).',
   '8 non-motorcycle sliver boxes found by the label audit are removed.',
   'Validation images and labels are untouched; models, folds, epochs and inference are identical to v1.'])
FIGURE('v2_copy_paste_examples.png', 'v2 copy-paste examples (red = pasted rare vehicle)')
P('A first A/B test of v2 on folds 1-2 (before the reuse cap was added) was kept only as evidence for the decision; '
  'all five v2 fold models used for the submission were trained with the same final code.')
H('3.4 Inference, ensemble and post-processing', 2)
B(['Confidence threshold 0.001: low-confidence boxes only extend the precision-recall curve and never '
   'reduce AP, so a very low threshold gives the highest mAP.',
   'Test-time augmentation (flips + scales) for every model.',
   '**Weighted Boxes Fusion (WBF)** of the fold models: boxes of the same class from different models that overlap '
   'are merged into one box whose coordinates are the confidence-weighted average; the score is averaged over the '
   'models, so a box found by only one model gets a low score. The first submission used IoU 0.55 / avg; tuning on the '
   'OOF predictions of folds 1-2 showed that 0.55 also merges neighbouring boxes of the same model, so the final '
   'setting is **IoU 0.7, conf_type box_and_model_avg** (Section 4.3). Up to 300 boxes per image are kept.',
   '**Final ensemble:** 10 models = v1 folds 0-4 + v2 folds 0-4 (experiments/ensemble_wbf.py).',
   'Boxes are written in absolute pixel coordinates (x1, y1, x2, y2); image_id is the test file name; images without '
   'detections get no row and no dummy boxes are added (Data-page rule).'])
H('3.5 Evaluation', 2)
P('Both notebooks re-implement the Kaggle metric with pycocotools (COCOeval, IoU = 0.50, area = all, maxDets = 100) '
  'on held-out cameras, using the raw labels. Ultralytics\' own mAP50 is a bit higher because it uses the cleaned labels.')

# ---------------------------------------------------------------- 4 results
H('Chapter 4: Results')
H('4.1 Experiments with YOLO26s (validation = 4 held-out cameras)', 2)
rows = [[name, e['desc'], fmt(e.get('val_map50')), fmt(e.get('kaggle_public'))]
        for name, e in RES['experiments'].items() if not name.startswith('yolo26m')]
TABLE(['Run', 'Description', 'Val mAP@50', 'Kaggle public'], rows, 'YOLO26s experiments', widths=[3.6, 8, 2.4, 2.4])
P('Higher resolution (640 -> 960) and TTA each add about +0.03. The YOLO26s results below are for the best '
  'single-split model (960 + TTA).')
e = RES['experiments']['yolo26s_960_tta']
TABLE(['Class'] + CLASSES, [['AP@50'] + [f"{e['val_ap'][c]:.3f}" for c in CLASSES]], 'Per-class AP@50, YOLO26s 960 + TTA, 4 validation cameras')
FIGURE(RES['figs']['curves'], 'YOLO26s 960: training curves (losses and validation mAP)')
FIGURE(RES['figs']['confusion'], 'YOLO26s 960: normalised confusion matrix on the validation cameras', width=13)
FIGURE(RES['figs']['gt_pred'], 'YOLO26s 960: ground truth (left) vs prediction with confidence > 0.3 (right)')

H('4.2 YOLO26m v1: camera 5-fold cross-validation', 2)
TABLE(['Fold', 'Validation cameras', 'Val mAP@50', 'Epochs run', 'Best epoch', 'Hours'],
      [[f['fold'], ', '.join(f['val_cams']), f"{f['val_map50']:.4f}", f['epochs_run'], f['best_epoch'], f['hours']] for f in CVR['folds']]
      + [['OOF', 'all 15 cameras (2,991 images)', f"{CVR['oof_map50']:.4f}", '', '', f"{sum(f['hours'] for f in CVR['folds']):.1f}"]],
      'YOLO26m 960 + TTA: per-fold and out-of-fold (OOF) validation mAP@50', widths=[1.3, 4.6, 2.3, 2.2, 2.2, 1.6],
      bold_rows=[len(CVR['folds'])])
FIGURE('cv_fold_curves.png', 'YOLO26m: validation mAP50 per epoch for each fold')
TABLE(['Class', 'YOLO26s (4 cams)', 'YOLO26m CV (same 4 cams)', 'YOLO26m CV (15 cams)'],
      [[c, f"{CVR['old4_s_ap'][c]:.3f}", f"{CVR['old4_m_ap'][c]:.3f}", f"{CVR['oof_ap'][c]:.3f}"] for c in CLASSES]
      + [['mAP@50', f"{CVR['old4_s']:.3f}", f"{CVR['old4_m']:.3f}", f"{CVR['oof_map50']:.3f}"]],
      'Per-class AP@50 on cameras that were not used for training', widths=[3, 4, 4.5, 4.5], bold_rows=[8])
FIGURE('cv_per_class.png', 'Per-class AP@50: YOLO26s vs YOLO26m on unseen cameras')
P(f"On the same four cameras YOLO26m is better for 6 of 8 classes (e.g. Van {CVR['old4_s_ap']['Van']:.2f} -> "
  f"{CVR['old4_m_ap']['Van']:.2f}, Car {CVR['old4_s_ap']['Car']:.2f} -> {CVR['old4_m_ap']['Car']:.2f}). "
  f"The OOF score over all 15 cameras ({CVR['oof_map50']:.3f}) is lower mainly because of Songthaew "
  f"(AP {CVR['oof_ap']['Songthaew']:.3f}). 91 of its 117 boxes come from camera 1426. In fold 0 (1426 in validation) "
  'only 26 Songthaew boxes are left for training; but even in fold 1, where 1426 is in the training set, Songthaew on '
  'camera 1066 is not detected (AP 0.004, 13 boxes). The model ties Songthaew to the close-up view of camera 1426 and '
  'does not transfer it to other cameras (see Chapter 5).')
FIGURE('cv_fold1_confusion.png', 'YOLO26m fold 1: normalised confusion matrix on its validation cameras (1066, 182, 222)', width=13)
H('4.3 YOLO26m v2 and the v1 + v2 ensemble', 2)
TABLE(['Fold', 'v1', 'v2', 'v1 (WBF 0.7)', 'v1 + v2 (WBF 0.7)'],
      [[r['fold'], f"{r['v1']:.4f}", f"{r['v2']:.4f}", f"{r['v1 via WBF0.7']:.4f}", f"{r['v1+v2 WBF0.7']:.4f}"] for r in V12['per_fold']]
      + [['mean', f"{sum(r['v1'] for r in V12['per_fold']) / 5:.4f}", f"{sum(r['v2'] for r in V12['per_fold']) / 5:.4f}",
          f"{sum(r['v1 via WBF0.7'] for r in V12['per_fold']) / 5:.4f}", f"{sum(r['v1+v2 WBF0.7'] for r in V12['per_fold']) / 5:.4f}"]],
      'Validation mAP@50 per fold (each model never saw the validation cameras of its fold)', widths=[1.5, 2.6, 2.6, 3.2, 3.6], bold_rows=[5])
TABLE(['Class', 'v1', 'v2', 'v1 + v2'],
      [[c, f"{V12['oof_v1w']['ap'][c]:.3f}", f"{V12['oof_v2']['ap'][c]:.3f}", f"{V12['oof_v12']['ap'][c]:.3f}"] for c in CLASSES]
      + [['OOF mAP@50', f"{V12['oof_v1w']['map50']:.4f}", f"{V12['oof_v2']['map50']:.4f}", f"{V12['oof_v12']['map50']:.4f}"]],
      'Out-of-fold per-class AP@50 over all 15 cameras (WBF IoU 0.7)', widths=[3.5, 3, 3, 3], bold_rows=[8])
P(f"v2 alone is on par with v1 (mean fold mAP {sum(r['v2'] for r in V12['per_fold']) / 5:.3f} vs {sum(r['v1'] for r in V12['per_fold']) / 5:.3f}): "
  'copy-paste improves Tuktuk (0.48 -> 0.52) but slightly hurts Bus and Motorcycle, and does not fix Songthaew. '
  f"Fusing v1 and v2, however, improves **every fold** (+0.03 on average) and the OOF mAP from {V12['oof_v1w']['map50']:.3f} "
  f"to {V12['oof_v12']['map50']:.3f}: the two runs make different errors, so the ensemble gains on Tuktuk (0.57), Van, Pickup and Truck.")
H('WBF parameter tuning (OOF, folds 1-2)', 3)
g = WBFT.groupby(['iou', 'conf_type'])[['v1_only', 'v1+v2']].mean().reset_index()
TABLE(['IoU threshold', 'conf_type', 'v1 alone', 'v1 + v2'],
      [[r['iou'], r['conf_type'], f"{r['v1_only']:.4f}", f"{r['v1+v2']:.4f}"] for _, r in g.iterrows()],
      'Mean mAP@50 of folds 1-2 for different WBF settings', widths=[3, 4.5, 3, 3])
P('IoU 0.55 lowers even a single model (0.511 -> 0.496) because it merges neighbouring vehicles of the same class; '
  'IoU 0.7 with box_and_model_avg is the best setting and was used for the final submissions.')
FIGURE('final_ensemble_test_pred.png', 'Final v1 + v2 ensemble predictions (confidence > 0.3) on the unseen test cameras')

H('4.4 Kaggle submissions', 2)
TABLE(['Submitted file', 'Model / image_id format', 'Public mAP@50'], RES['kaggle_history'], 'Kaggle submission history', widths=[5.4, 7.6, 2.7],
      bold_rows=[len(RES['kaggle_history']) - 1])
P('The first two submissions scored 0.000 although the model was the same: they used the long ids copied from '
  'sample_submission.csv, which do not exist in the solution file. Re-submitting the same predictions with the test '
  'file names gave 0.516, close to the validation score of the same model (0.525), which confirms that '
  'camera-held-out validation is a reliable estimate of the leaderboard. Training YOLO26s on all 15 cameras gave '
  '0.548 and the YOLO26m v1 5-fold WBF ensemble 0.605. Re-fusing the same v1 predictions with the tuned WBF setting '
  'gave 0.614 without any training, and the **v1 + v2 ensemble of 10 models gave 0.630**. The public scores follow '
  'the OOF scores (v2 alone < v1 < v1 + v2), so the model choice was made on cross-validation, not on the leaderboard.')
P(f"Final submission file: {final.get('file', '-')}. Public leaderboard mAP@50: {fmt(final.get('kaggle_public'))}.")
P('**Reproducibility check.** Re-running cv_kfold_kaggle.ipynb locally (RTX 3050 Ti) with the submitted fold weights '
  'reproduced every fold score within 0.0002 of the Kaggle run (e.g. fold 1: 0.5198 on Kaggle vs 0.5200 locally) and '
  'the ensemble file differs only by floating-point noise (6,948 vs 6,949 boxes with confidence > 0.3).')
P('Model weights of the final ensemble: weights/yolo26m_960_cv_fold{0..4}.pt (v1) and weights/yolo26m_960_cv_v2_fold{0..4}.pt (v2).')
FIGURE('kaggle_submissions_screenshot.png', 'Kaggle submission page (06/10/2026): the two selected final submissions are v1 + v2 (0.62963) and v1 with tuned WBF (0.61367)')

# ---------------------------------------------------------------- 5 discussion
H('Chapter 5: Discussion')
for title, items in RES['discussion'].items():
    H(title, 2)
    B(items)

# ---------------------------------------------------------------- 6 conclusion
H('Chapter 6: Conclusion')
for para in RES['conclusion']:
    P(para)

H('Appendix: How to reproduce', 1)
B(['conda create -n dsde python=3.12; pip install torch==2.8.0 torchvision==0.23.0 --index-url '
   'https://download.pytorch.org/whl/cu126; pip install -r requirements.txt',
   'kaggle competitions download -c 2110531-dsde-2026-1 -p data/raw, then unzip into data/raw',
   'Open vehicle_detection.ipynb from the midterm/ folder and run all cells. The settings are in the Config cell '
   '(or set environment variables, e.g. MODEL, IMGSZ, EPOCHS, FULL_TRAIN=1).',
   'YOLO26s submission without training: SKIP_TRAIN=1 WEIGHTS=<submitted .pt> (and the same IMGSZ / TTA as in the '
   'submission .json file).',
   'Final YOLO26m ensemble: run cv_kfold_kaggle.ipynb on a Kaggle Notebook (GPU T4 x2, attach the competition data) '
   'or locally. To reproduce without training, copy weights/yolo26m_960_cv_fold{k}.pt to cv_out/fold{k}/best.pt; '
   'the notebook then only predicts and writes cv_out/fold{k}/test_pred.csv (PRED_BATCH=8 on a 4 GB GPU). Same for v2 '
   '(cv_kfold_kaggle_v2.ipynb, weights/yolo26m_960_cv_v2_fold{k}.pt -> cv_out_v2/fold{k}/best.pt).',
   'Final CSV: python experiments/ensemble_wbf.py submissions/yolo26m_960_v1v2_10models_wbf07.csv 0.7 box_and_model_avg '
   'experiments/kaggle_cv_run/cv_out experiments/kaggle_cv_run_v2/cv_out_v2'])

out = ROOT / 'report' / 'midterm_report.docx'
doc.save(out)
print('saved', out)
try:
    from docx2pdf import convert
    convert(str(out), str(out.with_suffix('.pdf')))
    print('saved', out.with_suffix('.pdf'))
except Exception as e:  # Word not available
    print('PDF conversion skipped:', e)
